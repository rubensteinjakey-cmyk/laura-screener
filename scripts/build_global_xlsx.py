"""Build the global screen workbook from output/global/*.csv (scored companies only)."""
import csv
import re
import sys
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = "/home/user/laura-screener"
OUT = sys.argv[1]

rows = list(csv.DictReader(open(f"{ROOT}/output/global/results_global.csv", encoding="utf-8")))
companies = list(csv.DictReader(open(f"{ROOT}/companies_global.csv", encoding="utf-8-sig")))
order = {c["ticker"]: i for i, c in enumerate(companies)}
scored = [r for r in rows if r["final_score"]]
dropped = [r for r in rows if not r["final_score"]]
country_order = list(OrderedDict.fromkeys(c["country"] for c in companies))
scored.sort(key=lambda r: (country_order.index(r["country"]), r["sector"],
                           int(r["country_sector_rank"]), r["company_name"]))
evidence = list(csv.DictReader(open(f"{ROOT}/output/global/evidence_global.csv", encoding="utf-8")))
names = {r["ticker"]: r for r in rows}

FONT = "Arial"
H_FILL = PatternFill("solid", fgColor="1F3864")
H_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BASE = Font(name=FONT, size=10)
INPUT = Font(name=FONT, size=10, color="0000FF")
BOLD = Font(name=FONT, size=10, bold=True)
PASS_FILL = PatternFill("solid", fgColor="E2EFDA")
FAIL_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
WRAP = Alignment(wrap_text=True, vertical="top")
TOP = Alignment(vertical="top")


def num(x):
    try:
        f = float(x)
        return int(f) if f.is_integer() else f
    except (TypeError, ValueError):
        return x


def header(ws, row, cols, widths):
    for i, (name, w) in enumerate(zip(cols, widths), 1):
        c = ws.cell(row=row, column=i, value=name)
        c.font, c.fill = H_FONT, H_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[row].height = 30


def filing_parts(r):
    m = re.match(r"(.*?) (\d{4})? ?\((.*?)\) ?(.*)$", r["filing_used"])
    if not m:
        return r["filing_used"], "", r["source_site"], ""
    return m.group(1), m.group(2) or "", m.group(3), m.group(4)


wb = Workbook()

# ------------------------------------------------------------ All scored
ws = wb.active
ws.title = "All scored"
cols = ["Country", "Sector", "Ticker", "Company", "Creativity (0-5)", "Community (0-5)",
        "Cross-cultural (0-5)", "Entrepreneurship (0-5)", "Alignment score (0-100)",
        "Red flag deduction", "Final score", "Country-sector rank", "Values result",
        "Red flags", "Third component (TBD)", "Thin-coverage themes", "Score source",
        "Document", "Fiscal year", "Source site", "Report link", "Language", "Words read",
        "Notes", "Flags for review"]
widths = [12, 22, 12, 34, 10, 10, 11, 13, 11, 10, 9, 10, 9, 12, 11, 18, 30, 15, 8, 30, 40,
          11, 9, 60, 50]
header(ws, 1, cols, widths)
first, last = 2, 1 + len(scored)
for i, r in enumerate(scored, start=first):
    doc, fy, site, link = filing_parts(r)
    thin = [t.replace("_", "-") for t in ("creativity", "community", "cross_cultural",
                                          "entrepreneurship")
            if r[f"{t}_coverage"] == "insufficient"]
    vals = [r["country"], r["sector"], r["ticker"], r["company_name"],
            num(r["creativity"]), num(r["community"]), num(r["cross_cultural"]),
            num(r["entrepreneurship"]),
            f"=E{i}/5*35+F{i}/5*30+G{i}/5*20+H{i}/5*15",
            num(r["red_flag_score"]),
            f"=MAX(0,I{i}-J{i})",
            f'=COUNTIFS($A${first}:$A${last},A{i},$B${first}:$B${last},B{i},'
            f'$K${first}:$K${last},">"&K{i})+1',
            f'=IF(L{i}<=CEILING(COUNTIFS($A${first}:$A${last},A{i},$B${first}:$B${last},B{i})/2,1),'
            f'"Pass","Fail")',
            r["red_flags"] if r["red_flags"] != "none" else "",
            "", ", ".join(thin), r["score_source"], doc, fy, site, link,
            "Non-English" if r["non_english_source"] else "English",
            num(r["words_sent"]), r["values_notes"], r["flags"].replace(" | ", "\n")]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font = INPUT if j in (5, 6, 7, 8, 10) else BASE
        c.alignment = WRAP if j in (4, 24, 25) else TOP
        c.border = Border(bottom=THIN)
    ws.cell(row=i, column=9).number_format = "0.0"
    ws.cell(row=i, column=11).number_format = "0.0"
    ws.cell(row=i, column=23).number_format = "#,##0"
    if link:
        ws.cell(row=i, column=21).hyperlink = link
        ws.cell(row=i, column=21).font = Font(name=FONT, size=10, color="0563C1", underline="single")
ws.freeze_panes = "E2"
ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{last}"
from openpyxl.formatting.rule import FormulaRule
ws.conditional_formatting.add(f"M{first}:M{last}",
                              FormulaRule(formula=[f'M{first}="Pass"'], fill=PASS_FILL))
ws.conditional_formatting.add(f"M{first}:M{last}",
                              FormulaRule(formula=[f'M{first}="Fail"'], fill=FAIL_FILL))
note_row = last + 2
notes = [
    "Blue = inputs from the AI scoring after the code's quote checks (theme scores 0-5; red flag "
    "deduction = 10 points per verified red flag). Black = formulas.",
    "Alignment = Creativity/5×35 + Community/5×30 + Cross-cultural/5×20 + Entrepreneurship/5×15. "
    "Final = Alignment − red flag deduction (not below 0).",
    "Rank = 1 + companies in the same country-sector with a higher final score (ties share a rank). "
    "Pass = rank within the top half: rank ≤ ceil(n/2), n = scored companies in the country-sector.",
    "Decimal theme scores (e.g. 2.5) come from the half-weight rule: standalone sustainability-report "
    "evidence counts at half weight. Third component: not decided yet (blank).",
]
for k, t in enumerate(notes):
    c = ws.cell(row=note_row + k, column=1, value=t)
    c.font = Font(name=FONT, size=9, italic=True)

# ------------------------------------------------------------ Summary
sm = wb.create_sheet("Summary", 0)
sm["A1"] = "Rooted Capital global sleeve: values screen (scored companies)"
sm["A1"].font = Font(name=FONT, size=14, bold=True)
sm["A2"] = (f"As of Oct 23, 2026 · {len(scored)} companies scored in 14 country-sectors · "
            f"{len(dropped)} dropped by the team (see 'Dropped') · values filter only")
sm["A2"].font = Font(name=FONT, size=10, italic=True)
groups = list(OrderedDict.fromkeys((r["country"], r["sector"]) for r in scored))
hdr = ["Country", "Sector", "Companies scored", "Pass", "Fail", "Top company", "Top final score",
       "Average final score"]
header(sm, 4, hdr, [13, 24, 11, 8, 8, 40, 11, 11])
A = f"'All scored'!$A${first}:$A${last}"
B = f"'All scored'!$B${first}:$B${last}"
K = f"'All scored'!$K${first}:$K${last}"
L = f"'All scored'!$L${first}:$L${last}"
M = f"'All scored'!$M${first}:$M${last}"
D = f"'All scored'!$D${first}:$D${last}"
for i, (country, sector) in enumerate(groups, start=5):
    sm.cell(row=i, column=1, value=country)
    sm.cell(row=i, column=2, value=sector)
    sm.cell(row=i, column=3, value=f"=COUNTIFS({A},A{i},{B},B{i})")
    sm.cell(row=i, column=4, value=f'=COUNTIFS({A},A{i},{B},B{i},{M},"Pass")')
    sm.cell(row=i, column=5, value=f'=COUNTIFS({A},A{i},{B},B{i},{M},"Fail")')
    # first company (in sheet order) of this group with rank 1
    sm.cell(row=i, column=6, value=(
        f"=INDEX({D},MATCH(1,INDEX(({A}=A{i})*({B}=B{i})*({L}=1),0),0))"))
    sm.cell(row=i, column=7, value=f"=SUMPRODUCT(MAX(({A}=A{i})*({B}=B{i})*{K}))")
    sm.cell(row=i, column=8, value=f"=AVERAGEIFS({K},{A},A{i},{B},B{i})")
    for j in range(1, 9):
        c = sm.cell(row=i, column=j)
        c.font = BASE
        c.border = Border(bottom=THIN)
    sm.cell(row=i, column=7).number_format = "0.0"
    sm.cell(row=i, column=8).number_format = "0.0"
tot = 5 + len(groups)
sm.cell(row=tot, column=1, value="Total").font = BOLD
for col in (3, 4, 5):
    L_ = get_column_letter(col)
    c = sm.cell(row=tot, column=col, value=f"=SUM({L_}5:{L_}{tot - 1})")
    c.font = BOLD
c = sm.cell(row=tot, column=8, value=f"=AVERAGE({K})")
c.font, c.number_format = BOLD, "0.0"
sm.cell(row=tot + 2, column=1, value=(
    "Pass = top half of the country-sector by final score, ties included. Ties for the top spot "
    "show the first company listed; see 'All scored' for full ranks.")).font = Font(
    name=FONT, size=9, italic=True)
sm.cell(row=tot + 3, column=1, value=(
    "Sources: fiscal-2025 annual reports from official exchange filing sites (Nasdaq Copenhagen, "
    "MOPS, SGX) or company websites; 20-Fs from SEC EDGAR for Genmab, Alcon and Telkom.")).font = \
    Font(name=FONT, size=9, italic=True)
sm.freeze_panes = "A5"

# ------------------------------------------------------------ Pass list
pl = wb.create_sheet("Pass list", 1)
header(pl, 1, ["Country", "Sector", "Rank", "Ticker", "Company", "Final score",
               "Strongest theme"], [13, 24, 7, 12, 40, 10, 70])
k = 2
for r in scored:
    if r["values_result"] != "Pass":
        continue
    best = r["values_notes"].split("; ")[1] if "; " in r["values_notes"] else ""
    for j, v in enumerate([r["country"], r["sector"], int(r["country_sector_rank"]), r["ticker"],
                           r["company_name"], num(r["final_score"]), best], 1):
        c = pl.cell(row=k, column=j, value=v)
        c.font = BASE
        c.alignment = TOP
        c.border = Border(bottom=THIN)
    pl.cell(row=k, column=6).number_format = "0.0"
    k += 1
pl.cell(row=k + 1, column=1, value=(
    "Snapshot of the 'Values result' column on Oct 23, 2026 (sorted by country, sector, rank).")
        ).font = Font(name=FONT, size=9, italic=True)
pl.freeze_panes = "A2"
pl.auto_filter.ref = f"A1:G{k - 1}"

# ------------------------------------------------------------ Evidence
ev = wb.create_sheet("Evidence quotes")
header(ev, 1, ["Country", "Sector", "Ticker", "Company", "Type", "Themes / flag", "Activity",
               "Quote (word for word, checked by code)", "English translation", "Section",
               "Source"], [12, 20, 11, 28, 13, 26, 40, 70, 50, 26, 14])
kinds = {"evidence": "Evidence"}
k = 2
scored_t = {r["ticker"] for r in scored}
ev_rows = [e for e in evidence if e["ticker"] in scored_t and e["verified"] == "True"]
ev_rows.sort(key=lambda e: (order.get(e["ticker"], 999), e["kind"] != "evidence", e["id"]))
for e in ev_rows:
    r = names[e["ticker"]]
    if e["kind"] == "evidence":
        typ = "Evidence"
    elif "found by: internal" in e["themes_or_flag"]:
        typ = "Internal finding (no deduction)"
    else:
        typ = "Red flag"
    for j, v in enumerate([r["country"], r["sector"], e["ticker"], r["company_name"], typ,
                           e["themes_or_flag"], e["activity"], e["quote"], e["quote_en"],
                           e["section"], e["source_doc"]], 1):
        c = ev.cell(row=k, column=j, value=v)
        c.font = BASE
        c.alignment = WRAP if j in (6, 7, 8, 9, 10) else TOP
        c.border = Border(bottom=THIN)
    k += 1
ev.freeze_panes = "E2"
ev.auto_filter.ref = f"A1:K{k - 1}"

# ------------------------------------------------------------ Dropped
dr = wb.create_sheet("Dropped")
header(dr, 1, ["Country", "Sector", "Ticker", "Company", "Why it was dropped"],
       [12, 24, 12, 34, 90])
for k, r in enumerate(dropped, start=2):
    for j, v in enumerate([r["country"], r["sector"], r["ticker"], r["company_name"],
                           r["values_notes"]], 1):
        c = dr.cell(row=k, column=j, value=v)
        c.font = BASE
        c.alignment = WRAP if j == 5 else TOP

wb.save(OUT)
print(f"saved {OUT}: {len(scored)} scored, {len(ev_rows)} quotes, {len(dropped)} dropped")
